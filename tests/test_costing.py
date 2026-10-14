"""The cost engine must reproduce the original 'Straight' sheet exactly."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from hydroquote import CostInputs, PartSpec, compute_cost

SHEET = Path(__file__).resolve().parent.parent / "samples" / "cost_sheet.xlsx"


def _rows():
    if not SHEET.exists():  # samples are not in the repo; see README
        return
    ws = load_workbook(SHEET, data_only=True)["Straight"]
    for r in range(2, ws.max_row + 1):
        v = {c: ws[f"{c}{r}"].value for c in
             ("B", "C", "D", "E", "F", "G", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T",
              "U", "V", "W", "X", "Y", "Z", "AA")}
        if v["B"]:
            yield v


@pytest.mark.skipif(not SHEET.exists(), reason="samples/cost_sheet.xlsx not present")
@pytest.mark.parametrize("row", list(_rows()) or [None], ids=lambda r: r["B"].strip() if r else "none")
def test_matches_original_sheet(row):
    part = PartSpec(description=row["B"], material=row["D"], shape=row["E"],
                    bar_size=row["F"], overall_length=row["G"])
    cost_in = CostInputs(quantity=row["C"], raw_material_rate=row["K"], caps=row["Q"],
                         crimp_n=row["R"], nut_red_n=row["S"], ferrule_wn=row["T"], cn=row["U"],
                         was=row["V"], o_ring=row["W"], sr_ring=row["X"], margin=row["Z"])
    c = compute_cost(part, cost_in, row["N"])
    assert c.rm_weight_kg == pytest.approx(row["J"])
    assert c.rm_cost == pytest.approx(row["L"])
    assert c.cutting == pytest.approx(row["M"])
    assert c.machining == pytest.approx(row["O"])
    assert c.plating == pytest.approx(row["P"])
    assert c.net == pytest.approx(row["Y"])
    assert c.selling_price == pytest.approx(row["AA"])


def test_nut_and_ferrule_optional():
    part = PartSpec(description="x", bar_size=19, overall_length=30)
    base = compute_cost(part, CostInputs(nut_red_n=20, ferrule_wn=10), 4.5)
    incl = compute_cost(part, CostInputs(nut_red_n=20, ferrule_wn=10), 4.5, include_nut_ferrule=True)
    assert incl.net == pytest.approx(base.net + 30)
