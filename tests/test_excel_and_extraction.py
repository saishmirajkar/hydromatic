import json
from io import BytesIO
from pathlib import Path

import anthropic
import httpx2
import pytest
from openpyxl import load_workbook

from hydroquote import CostInputs, QuoteLine, load_calibration, load_machine, load_reference_parts
from hydroquote import estimate_with_calibration
from hydroquote.excel_io import build_straight_workbook, read_straight_rows
from hydroquote.extraction import MODEL, extract_part_from_drawing

ROOT = Path(__file__).resolve().parent.parent
SHEET = ROOT / "samples" / "cost_sheet.xlsx"
DRAWING = ROOT / "samples" / "drawings" / "G_10_S.pdf"


def test_export_has_straight_formulas_and_breakdown():
    machine, cal = load_machine(), load_calibration()
    refs = load_reference_parts()[:2]
    ests = [estimate_with_calibration(r.part, machine, cal) for r in refs]
    lines = [QuoteLine(part=r.part, cost=CostInputs(caps=4), cycle_time_min=e.cycle_time_rounded_min)
             for r, e in zip(refs, ests)]
    wb = load_workbook(BytesIO(build_straight_workbook(lines, ests)))
    ws = wb["Straight"]
    assert ws["B1"].value == "Description" and ws["N1"].value == "Cycle Time"
    assert ws["J2"].value == "=0.8666*7.85*H2*H2*I2/1000"
    assert ws["Y3"].value == "=+L3+M3+O3+P3+Q3+R3+U3+V3+W3+X3"
    assert ws["N2"].value == lines[0].cycle_time_min
    assert "TL breakdown" in wb.sheetnames
    rows = read_straight_rows(build_straight_workbook(lines, ests))
    assert [r["description"] for r in rows] == [r.part.description for r in refs]


@pytest.mark.skipif(not SHEET.exists(), reason="samples/cost_sheet.xlsx not present")
def test_read_original_sheet():
    rows = read_straight_rows(SHEET.read_bytes())
    assert len(rows) == 8
    assert rows[0]["cycle_time_min"] == 5


@pytest.mark.skipif(not DRAWING.exists(), reason="sample drawing not present")
def test_extraction_request_shape():
    """Check the Claude request without calling the API."""
    ref = json.loads((ROOT / "data" / "reference_parts" / "G_10_S.json").read_text())["part"]
    seen = {}

    def handler(request):
        seen["body"] = json.loads(request.content)
        seen["beta"] = request.headers.get("anthropic-beta")
        return httpx2.Response(200, json={
            "id": "msg_test", "type": "message", "role": "assistant", "model": MODEL,
            "content": [{"type": "text", "text": json.dumps(ref)}],
            "stop_reason": "end_turn", "stop_sequence": None,
            "usage": {"input_tokens": 1, "output_tokens": 1}})

    client = anthropic.Anthropic(api_key="test", http_client=httpx2.Client(
        transport=httpx2.MockTransport(handler)))
    pdf = DRAWING.read_bytes()
    part = extract_part_from_drawing(pdf, "G_10_S.pdf", client)

    assert part.description == "G 10 - S" and len(part.features) == 15
    body = seen["body"]
    assert body["model"] == MODEL
    assert body["output_config"]["format"]["type"] == "json_schema"
    content = body["messages"][0]["content"]
    assert content[0]["type"] == "document"
    assert "M18X1.5" in content[1]["text"]  # vector text from the PDF is passed along
    assert "server-side-fallback" in seen["beta"]
